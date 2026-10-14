"""Reading import files (CSV or XLSX) and building the template."""

import csv
import io
from collections.abc import Iterator
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from app.core.errors import InvalidInput

COLUMNS = [
    "employee_code", "first_name", "last_name", "email", "job_title", "job_level",
    "department", "country_code", "city", "employment_type", "fte", "hire_date", "gender",
    "salary", "currency", "salary_effective_date", "salary_reason", "salary_note",
]  # fmt: skip
REQUIRED = [
    "employee_code", "first_name", "last_name", "email", "job_title", "job_level",
    "department", "country_code", "employment_type", "hire_date", "salary",
]  # fmt: skip

HELP = {
    "employee_code": "Required. Matches existing people (e.g. ACM-004521); new codes add people",
    "job_level": "Required. 1 to 6",
    "department": "Required. Must match a department name in the app",
    "country_code": "Required. IN, US, CA, SG, BR or AU",
    "employment_type": "Required. full_time, part_time or contractor",
    "fte": "Optional. Part-time share, e.g. 0.5 (default 1)",
    "hire_date": "Required. YYYY-MM-DD",
    "gender": "Optional. female, male, non_binary or undisclosed",
    "salary": "Required. Yearly base salary, e.g. 1840000",
    "currency": "Optional. Default: the employee's current one (new people: the country's)",
    "salary_effective_date": "Optional. YYYY-MM-DD; when a salary changes (default: today)",
    "salary_reason": "Optional. promotion, merit, market_adjustment, correction or other",
}

EXAMPLE = [
    "ACM-010001", "Asha", "Verma", "asha.verma@acme.example", "Software Engineer", 3,
    "Engineering", "IN", "Bengaluru", "full_time", 1, "2024-04-01", "", 1840000, "INR",
    "2024-04-01", "", "",
]  # fmt: skip

XLSX_MAGIC = b"PK\x03\x04"


def file_kind(file_name: str | None, data: bytes) -> str:
    """'xlsx' or 'csv', judged by content; macro-enabled and old Excel files are refused."""
    name = (file_name or "").lower()
    if name.endswith((".xlsm", ".xls", ".xlsb")):
        raise InvalidInput("Save the file as .xlsx (without macros) or .csv and try again.")
    if data.startswith(XLSX_MAGIC):
        if b"vbaProject" in data:
            raise InvalidInput("Files with macros are not allowed. Save as .xlsx or .csv.")
        return "xlsx"
    try:
        data[:65536].decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvalidInput("Only .xlsx and .csv (UTF-8) files can be imported.") from exc
    if b"\x00" in data[:65536]:
        raise InvalidInput("Only .xlsx and .csv (UTF-8) files can be imported.")
    return "csv"


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(cells: list) -> list[str]:
    return [_cell(c).lower().replace(" ", "_") for c in cells]


def read_rows(kind: str, data: bytes) -> tuple[list[str], Iterator[tuple[int, dict[str, str]]]]:
    """Return the header and an iterator of (row number as shown in Excel, values).

    Blank rows are skipped. Raises InvalidInput if required columns are missing.
    """
    if kind == "csv":
        reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        rows = iter(reader)
    else:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = workbook.worksheets[0].iter_rows(values_only=True)
    try:
        header = _header(list(next(rows)))
    except StopIteration as exc:
        raise InvalidInput("The file is empty.") from exc
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise InvalidInput(f"Missing columns: {', '.join(missing)}. Start from the template.")

    def generate() -> Iterator[tuple[int, dict[str, str]]]:
        for number, cells in enumerate(rows, start=2):
            values = {name: _cell(v) for name, v in zip(header, cells, strict=False) if name}
            if any(values.values()):
                yield number, {c: values.get(c, "") for c in COLUMNS}

    return header, generate()


def template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(COLUMNS)
    sheet.append(EXAMPLE)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 18
    notes = workbook.create_sheet("How to fill")
    notes.append(["Column", "What to enter"])
    for name in COLUMNS:
        notes.append([name, HELP.get(name, "Required" if name in REQUIRED else "Optional")])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 80
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
