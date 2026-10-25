import asyncio
import csv
import io
import uuid
from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.dialects import postgresql

from app.models import Employee, ImportJob, SalaryRecord
from app.services import import_file
from app.services import imports as service
from app.services.import_file import COLUMNS, template_xlsx
from app.worker.tasks import process_next_job
from tests.conftest import count

API = "/api/v1/imports"
HEADER = ",".join(COLUMNS)


def row(**values) -> dict:
    base = {
        "employee_code": "ACM-100001", "first_name": "Asha", "last_name": "Verma",
        "email": "asha@acme.example", "job_title": "Software Engineer", "job_level": "3",
        "department": "Engineering", "country_code": "IN", "city": "Pune",
        "employment_type": "full_time", "fte": "", "hire_date": "2025-04-01", "gender": "",
        "salary": "1840000", "currency": "", "salary_effective_date": "", "salary_reason": "",
        "salary_note": "",
    }  # fmt: skip
    base.update(values)
    return base


def csv_file(*rows: dict) -> bytes:
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode()


def xlsx_file(*rows: list) -> bytes:
    workbook = Workbook()
    for values in rows:
        workbook.active.append(values)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def start(client, data: bytes, name="people.csv", key=None):
    headers = {"Idempotency-Key": key} if key else {}
    return client.post(API, files={"file": (name, data)}, headers=headers)


async def run_worker(session_factory, storage):
    async with session_factory() as db:
        while await process_next_job(db, storage):
            pass


@pytest.fixture
async def ctx(db):
    return lambda: service.load_context(db)


# ---------- row checks (no worker) ----------
async def test_valid_new_row_is_a_create(ctx):
    action, payload, errors = service.check_row(row(), await ctx())
    assert (action, errors) == ("create", [])
    assert payload["salary"] == {"amount": "1840000.00", "currency": "INR", "effective_date": "2025-04-01",
                                 "reason": "hire", "note": None}  # fmt: skip
    assert payload["values"]["fte"] == "1.000"


async def test_every_problem_in_a_row_is_reported(ctx):
    action, _, errors = service.check_row(
        row(email="nope", job_level="9", department="Space", country_code="FR", salary="-1",
            hire_date="01/04/2025", employment_type="intern", fte="2"),
        await ctx(),
    )  # fmt: skip
    assert action == "error"
    assert {e["field"] for e in errors} == {
        "email", "job_level", "department", "country_code", "salary", "hire_date",
        "employment_type", "fte",
    }  # fmt: skip


async def test_missing_required_values(ctx):
    _, _, errors = service.check_row(row(first_name="", salary="", employee_code=""), await ctx())
    assert {(e["field"], e["message"]) for e in errors} >= {
        ("first_name", "Required"), ("salary", "Required"), ("employee_code", "Required"),
    }  # fmt: skip


async def test_duplicates_inside_the_file(ctx):
    context = await ctx()
    assert service.check_row(row(), context)[0] == "create"
    action, _, errors = service.check_row(row(), context)
    assert action == "error"
    assert {e["field"] for e in errors} == {"employee_code", "email"}


async def test_existing_employee_rows(ctx, make_employee):
    e = await make_employee(country_code="IN", salary="1840000", email="asha@acme.example",
                      first_name="Asha", last_name="Verma", city="Pune", hire_date="2025-04-01")  # fmt: skip
    same = row(employee_code=e.employee_code)
    assert service.check_row(same, await ctx())[0] == "unchanged"
    # Blank optional cells leave stored values alone.
    assert service.check_row({**same, "city": ""}, await ctx())[0] == "unchanged"

    action, payload, _ = service.check_row(
        {**same, "job_level": "4", "salary": "2000000"}, await ctx()
    )
    assert action == "update"
    assert payload["values"] == {"job_level": 4}
    assert payload["salary"]["effective_date"] == "2026-10-14"  # default: today
    assert payload["salary"]["reason"] == "other"
    assert payload["version"] == e.version


async def test_blank_currency_keeps_the_employees_currency(ctx, make_employee):
    # Paid in USD while living in India: a blank currency cell must not mean INR.
    e = await make_employee(country_code="IN", salary="100000", currency="USD", email="asha@acme.example",
                      first_name="Asha", last_name="Verma", city="Pune", hire_date="2025-04-01")  # fmt: skip
    same = row(employee_code=e.employee_code, salary="100000")
    assert service.check_row(same, await ctx())[0] == "unchanged"
    action, payload, _ = service.check_row({**same, "salary": "110000"}, await ctx())
    assert action == "update"
    assert payload["salary"]["currency"] == "USD"


async def test_someone_not_started_yet_can_be_reimported(ctx, make_employee):
    e = await make_employee(country_code="IN", salary="1840000", email="asha@acme.example",
                      first_name="Asha", last_name="Verma", city="Pune", hire_date="2026-11-01")  # fmt: skip
    same = row(employee_code=e.employee_code, hire_date="2026-11-01")
    assert service.check_row(same, await ctx()) == ("unchanged", {"raw": same}, [])
    # A new salary without a date starts on their first day, not today.
    action, payload, _ = service.check_row({**same, "salary": "1900000"}, await ctx())
    assert action == "update"
    assert payload["salary"]["effective_date"] == "2026-11-01"


async def test_email_owned_by_someone_else(ctx, make_employee):
    await make_employee(email="asha@acme.example")
    _, _, errors = service.check_row(row(), await ctx())
    assert errors == [{"field": "email", "message": "Already used by another employee"}]


@pytest.mark.parametrize(
    ("values", "error"),
    [
        ({"first_name": "x" * 101}, {"field": "first_name", "message": "At most 100 characters"}),
        ({"currency": "EUR"}, {"field": "currency", "message": "Not a supported currency"}),
        ({"salary": "Infinity"}, {"field": "salary", "message": "Must be a number"}),
    ],
)
async def test_too_long_text_and_unknown_currency(ctx, values, error):
    action, _, errors = service.check_row(row(**values), await ctx())
    assert (action, errors) == ("error", [error])


async def test_no_salary_change_for_someone_who_has_left(ctx, make_employee, db):
    e = await make_employee(country_code="IN", salary="1840000", email="asha@acme.example",
                      first_name="Asha", last_name="Verma", city="Pune", hire_date="2025-04-01")  # fmt: skip
    e.status = "left"
    await db.commit()
    same = row(employee_code=e.employee_code)
    assert service.check_row(same, await ctx())[0] == "unchanged"
    _, _, errors = service.check_row({**same, "salary": "2000000"}, await ctx())
    assert errors == [
        {"field": "salary", "message": "This employee has left; salary changes aren't allowed"}
    ]


async def test_dates_must_be_sensible(ctx):
    _, _, errors = service.check_row(row(salary_effective_date="2024-01-01"), await ctx())
    assert errors[0]["field"] == "salary_effective_date"
    _, _, errors = service.check_row(row(hire_date="2028-01-01"), await ctx())
    assert {e["field"] for e in errors} == {"hire_date", "salary_effective_date"}


@pytest.mark.parametrize(
    ("cell", "text"),
    [
        (None, ""),
        (datetime(2025, 4, 1, 0, 0), "2025-04-01"),
        (date(2025, 4, 1), "2025-04-01"),
        (3.0, "3"),  # Excel keeps whole numbers as floats
        (0.5, "0.5"),
        ("  Pune ", "Pune"),
    ],
)
def test_spreadsheet_cells_are_read_as_text(cell, text):
    assert import_file._cell(cell) == text


def test_template_has_columns_example_and_help():
    workbook = load_workbook(io.BytesIO(template_xlsx()))
    sheet = workbook["Employees"]
    assert [c.value for c in sheet[1]] == COLUMNS
    assert sheet["A2"].value == "ACM-010001"
    assert workbook["How to fill"].max_row == len(COLUMNS) + 1


# ---------- full lifecycle through the API and the worker ----------
async def test_import_lifecycle_creates_and_updates(
    admin_client, make_employee, session_factory, storage, db
):
    existing = await make_employee(country_code="IN", salary="1500000", email="old@acme.example",
                             first_name="Old", last_name="Timer", hire_date="2020-01-01")  # fmt: skip
    untouched = await make_employee()
    data = csv_file(
        row(),
        row(employee_code=existing.employee_code, email="old@acme.example", first_name="Old",
            last_name="Timer", hire_date="2020-01-01", salary="1650000",
            salary_effective_date="2026-10-01", salary_reason="merit"),
    )  # fmt: skip
    job = start(admin_client, data).json()
    assert job["status"] == "queued"

    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job['id']}").json()
    assert job["status"] == "validated"
    assert (job["total_rows"], job["processed_rows"], job["create_count"], job["update_count"]) == (
        2,
        2,
        1,
        1,
    )
    rows = admin_client.get(f"{API}/{job['id']}/rows").json()
    assert [r["action"] for r in rows["items"]] == ["create", "update"]
    assert not list(storage.root.glob("imports/*"))  # raw file removed after checking

    assert (
        admin_client.post(f"{API}/{job['id']}/commit", json={}).json()["status"] == "commit_queued"
    )
    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job['id']}").json()
    assert (job["status"], job["created_count"], job["updated_count"]) == ("completed", 1, 1)

    new = await db.scalar(select(Employee).filter_by(employee_code="ACM-100001"))
    assert (new.current_salary, new.current_currency, str(new.current_salary_usd)) == (
        1840000,
        "INR",
        "22080.00",
    )
    await db.refresh(existing)  # the worker changed it in another session
    assert existing.current_salary == 1650000
    record = await db.scalar(
        select(SalaryRecord).filter_by(employee_id=existing.id, reason="merit")
    )
    assert record.import_job_id == job["id"]
    await db.refresh(untouched)
    assert untouched.version == 1  # people missing from the file never change


async def test_problems_block_all_or_nothing_but_correct_rows_can_be_saved(
    admin_client, session_factory, storage, db
):
    data = csv_file(row(), row(employee_code="ACM-100002", email="b@acme.example", salary="abc"))
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    assert admin_client.get(f"{API}/{job_id}").json()["error_count"] == 1

    blocked = admin_client.post(f"{API}/{job_id}/commit", json={"mode": "all_or_nothing"})
    assert blocked.status_code == 409

    report = list(csv.DictReader(io.StringIO(admin_client.get(f"{API}/{job_id}/errors.csv").text)))
    assert [(r["row"], r["salary"], r["problems"]) for r in report] == [
        ("3", "abc", "salary: Must be a number")
    ]
    errors = admin_client.get(f"{API}/{job_id}/rows", params={"action": "error"}).json()
    assert errors["total"] == 1 and errors["items"][0]["row_number"] == 3

    admin_client.post(f"{API}/{job_id}/commit", json={"mode": "valid_rows_only"})
    await run_worker(session_factory, storage)
    assert admin_client.get(f"{API}/{job_id}").json()["created_count"] == 1
    assert await count(db, Employee) == 1


def test_only_one_import_at_a_time_and_cancel(admin_client, session_factory, storage):
    first = start(admin_client, csv_file(row())).json()
    busy = start(admin_client, csv_file(row(employee_code="ACM-2")))
    assert busy.status_code == 409
    assert busy.json()["error"]["details"]["import_id"] == first["id"]
    cancelled = admin_client.post(f"{API}/{first['id']}/cancel").json()
    assert cancelled["status"] == "cancelled"
    assert admin_client.post(f"{API}/{first['id']}/cancel").status_code == 409
    assert admin_client.post(f"{API}/{first['id']}/commit", json={}).status_code == 409
    assert start(admin_client, csv_file(row(employee_code="ACM-2"))).status_code == 202


async def test_problem_report_blocks_spreadsheet_formulas(admin_client, session_factory, storage):
    bad = row(first_name='=HYPERLINK("http://evil.example","x")', city="@SUM(1)", salary="-5")
    job_id = start(admin_client, csv_file(bad)).json()["id"]
    await run_worker(session_factory, storage)
    report = next(csv.DictReader(io.StringIO(admin_client.get(f"{API}/{job_id}/errors.csv").text)))
    assert report["first_name"] == '\'=HYPERLINK("http://evil.example","x")'
    assert report["city"] == "'@SUM(1)"
    assert report["salary"] == "-5"  # a plain number stays as it is


async def test_two_uploads_at_once_still_open_only_one_import(
    admin_client, storage, db, monkeypatch
):
    """The database refuses a second open import even when both requests pass the check."""
    monkeypatch.setattr(service, "OPEN_IMPORT_STATUSES", ())  # both requests see no open import
    assert start(admin_client, csv_file(row())).status_code == 202
    files_before = set(storage.root.glob("imports/*"))
    second = start(admin_client, csv_file(row(employee_code="ACM-2")))
    assert second.status_code == 409
    assert await count(db, ImportJob) == 1
    assert set(storage.root.glob("imports/*")) == files_before  # the refused file is removed


async def test_same_upload_twice_starts_one_job(admin_client, db):
    key = str(uuid.uuid4())
    first, second = (
        start(admin_client, csv_file(row()), key=key),
        start(admin_client, csv_file(row()), key=key),
    )
    assert first.status_code == second.status_code == 202
    assert first.json()["id"] == second.json()["id"]
    assert await count(db, ImportJob) == 1


async def test_commit_twice_with_same_key_queues_once(admin_client, session_factory, storage):
    job_id = start(admin_client, csv_file(row())).json()["id"]
    await run_worker(session_factory, storage)
    key = str(uuid.uuid4())
    first = admin_client.post(f"{API}/{job_id}/commit", json={}, headers={"Idempotency-Key": key})
    second = admin_client.post(f"{API}/{job_id}/commit", json={}, headers={"Idempotency-Key": key})
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json()


async def test_reimporting_the_same_file_is_flagged(admin_client, session_factory, storage):
    data = csv_file(row())
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    admin_client.post(f"{API}/{job_id}/commit", json={})
    await run_worker(session_factory, storage)
    again = start(admin_client, data).json()
    assert again["already_imported_at"] is not None
    await run_worker(session_factory, storage)
    assert admin_client.get(f"{API}/{again['id']}").json()["unchanged_count"] == 1


async def test_employee_changed_after_check_stops_all_or_nothing(
    admin_client, make_employee, session_factory, storage, db
):
    e = await make_employee(email="x@acme.example", first_name="X", last_name="Y", country_code="IN",
                      salary="1000000", hire_date="2024-01-01")  # fmt: skip
    data = csv_file(row(), row(employee_code=e.employee_code, email="x@acme.example", first_name="X",
                               last_name="Y", hire_date="2024-01-01", salary="1100000"))  # fmt: skip
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    admin_client.patch(f"/api/v1/employees/{e.id}", json={"version": 1, "city": "Goa"})
    admin_client.post(f"{API}/{job_id}/commit", json={})
    await run_worker(session_factory, storage)

    job = admin_client.get(f"{API}/{job_id}").json()
    assert (
        job["status"] == "failed" and "The employee changed after the check" in job["error_message"]
    )
    assert await count(db, Employee) == 1  # the new row was rolled back too
    await db.refresh(e)
    assert e.current_salary == 1000000


async def test_code_or_email_taken_after_the_check(
    admin_client, make_employee, session_factory, storage, db
):
    data = csv_file(
        row(employee_code="ACM-000001"),
        row(employee_code="ACM-100002", email="b@acme.example"),
        row(employee_code="ACM-100003", email="c@acme.example"),
    )
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    assert admin_client.get(f"{API}/{job_id}").json()["create_count"] == 3
    # While the import waits for review, the UI adds ACM-000001 and someone with b@.
    await make_employee()
    await make_employee(email="b@acme.example")

    admin_client.post(f"{API}/{job_id}/commit", json={"mode": "valid_rows_only"})
    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job_id}").json()
    assert (job["status"], job["created_count"], job["error_count"]) == ("completed", 1, 2)
    errors = admin_client.get(f"{API}/{job_id}/rows", params={"action": "error"}).json()
    assert [(r["row_number"], r["errors"][0]["field"]) for r in errors["items"]] == [
        (2, "employee_code"),
        (3, "email"),
    ]
    assert await count(db, Employee, employee_code="ACM-100003") == 1


async def test_new_email_taken_after_the_check(
    admin_client, make_employee, session_factory, storage, db
):
    e = await make_employee(email="x@acme.example", first_name="X", last_name="Y", country_code="IN",
                      salary="1000000", hire_date="2024-01-01")  # fmt: skip
    data = csv_file(row(employee_code=e.employee_code, email="new@acme.example", first_name="X",
                        last_name="Y", hire_date="2024-01-01", salary="1000000"))  # fmt: skip
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    await make_employee(
        email="new@acme.example"
    )  # someone takes the address while it waits for review

    admin_client.post(f"{API}/{job_id}/commit", json={"mode": "valid_rows_only"})
    await run_worker(session_factory, storage)
    errors = admin_client.get(f"{API}/{job_id}/rows", params={"action": "error"}).json()
    assert errors["items"][0]["errors"] == [
        {"field": "email", "message": "Taken by another employee after the check"}
    ]
    await db.refresh(e)
    assert e.email == "x@acme.example"


async def test_code_taken_after_the_check_stops_all_or_nothing(
    admin_client, make_employee, session_factory, storage, db
):
    job_id = start(admin_client, csv_file(row(employee_code="ACM-000001"))).json()["id"]
    await run_worker(session_factory, storage)
    await make_employee()
    admin_client.post(f"{API}/{job_id}/commit", json={"mode": "all_or_nothing"})
    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job_id}").json()
    assert job["status"] == "failed"
    assert job["error_message"].startswith("Row 2, employee_code: Taken by an employee")
    assert await count(db, Employee) == 1


async def test_failure_in_the_middle_of_apply_saves_nothing(
    admin_client, session_factory, storage, db, monkeypatch
):
    data = csv_file(row(), row(employee_code="ACM-100002", email="b@acme.example"))
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    admin_client.post(f"{API}/{job_id}/commit", json={})

    real_create = service._create_from_row
    calls = []

    async def create_then_fail(*args):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("disk full")
        return await real_create(*args)

    monkeypatch.setattr(service, "_create_from_row", create_then_fail)
    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job_id}").json()
    assert job["status"] == "failed"
    assert await count(db, Employee) == 0


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("people.csv", b"first_name\nAsha\n", "Missing columns"),
        ("people.csv", HEADER.encode() + b"\n", "no data rows"),
        ("people.xlsx", xlsx_file(), "The file is empty"),
    ],
)
async def test_bad_files_fail_with_a_clear_message(
    admin_client, session_factory, storage, name, data, message
):
    job_id = start(admin_client, data, name=name).json()["id"]
    await run_worker(session_factory, storage)
    job = admin_client.get(f"{API}/{job_id}").json()
    assert job["status"] == "failed" and message in job["error_message"]
    assert not (storage.root / "imports" / job_id).exists()  # a failed job's file is not kept


async def test_row_limit(admin_client, session_factory, storage, monkeypatch):
    monkeypatch.setattr(service.get_settings(), "max_import_rows", 1)
    job_id = start(
        admin_client, csv_file(row(), row(employee_code="ACM-2", email="b@acme.example"))
    ).json()["id"]
    await run_worker(session_factory, storage)
    assert "limit is 1" in admin_client.get(f"{API}/{job_id}").json()["error_message"]


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("macro.xlsm", b"PK\x03\x04..."),
        ("binary.csv", b"\x00\x01\x02\xff\xfe"),
        ("old.xls", b"\xd0\xcf"),
        ("renamed.xlsx", b"PK\x03\x04 xl/vbaProject.bin"),
        ("nul.csv", b"employee_code\x00\n"),
    ],
)
def test_wrong_file_types_are_refused_at_upload(admin_client, name, data):
    assert start(admin_client, data, name=name).status_code == 422


async def test_xlsx_files_are_read(admin_client, session_factory, storage):
    # Excel hands back dates as datetimes and whole numbers as floats.
    typed = {**row(), "job_level": 3.0, "salary": 1840000.0, "hire_date": datetime(2025, 4, 1),
             "salary_effective_date": date(2025, 4, 1)}  # fmt: skip
    data = xlsx_file(COLUMNS, [typed[c] for c in COLUMNS])
    job_id = start(admin_client, data, name="people.xlsx").json()["id"]
    await run_worker(session_factory, storage)
    assert admin_client.get(f"{API}/{job_id}").json()["create_count"] == 1
    salary = admin_client.get(f"{API}/{job_id}/rows").json()["items"][0]["payload"]["salary"]
    assert (salary["amount"], salary["effective_date"]) == ("1840000.00", "2025-04-01")


@pytest.mark.parametrize(
    ("status", "quiet_minutes", "expected"),
    [
        ("validating", 16, "queued"),
        ("committing", None, "commit_queued"),  # no heartbeat at all
        ("committing", 1, "committing"),  # another worker is still on this one
    ],
)
async def test_interrupted_jobs_are_requeued(
    db, admin, frozen_clock, status, quiet_minutes, expected
):
    heartbeat = (
        None if quiet_minutes is None else frozen_clock.current - timedelta(minutes=quiet_minutes)
    )
    db.add(ImportJob(id="a", filename="a.csv", file_sha256="1", status=status,
                     created_by=admin.id, heartbeat_at=heartbeat))  # fmt: skip
    await db.commit()
    await service.reset_stuck_jobs(db)
    await service.reset_stuck_jobs(db)  # safe to re-run
    assert (await db.get(ImportJob, "a")).status == expected


async def test_claiming_and_checking_keep_the_heartbeat_fresh(
    admin_client, session_factory, storage, db, frozen_clock
):
    job_id = start(admin_client, csv_file(row())).json()["id"]
    await run_worker(session_factory, storage)
    db.expire_all()
    assert (await db.get(ImportJob, job_id)).heartbeat_at == frozen_clock.current


async def test_applying_keeps_the_heartbeat_fresh_while_its_transaction_is_open(
    admin_client, make_employee, session_factory, storage, frozen_clock, monkeypatch
):
    data = csv_file(
        row(employee_code="ACM-000001", email="a@acme.example"),
        row(employee_code="ACM-100002", email="b@acme.example"),
        row(employee_code="ACM-100003", email="c@acme.example"),
    )
    job_id = start(admin_client, data).json()["id"]
    await run_worker(session_factory, storage)
    await make_employee()  # row 2 is skipped as a problem before the other rows are applied
    admin_client.post(f"{API}/{job_id}/commit", json={"mode": "valid_rows_only"})

    monkeypatch.setattr(service, "CHUNK", 1)
    real_create = service._create_from_row
    seen = []

    async def create_and_look_from_outside(*args):
        async with session_factory() as other:  # what another worker sees mid-apply
            seen.append((await other.get(ImportJob, job_id)).heartbeat_at)
        frozen_clock.set(frozen_clock.current + timedelta(minutes=1))
        return await real_create(*args)

    monkeypatch.setattr(service, "_create_from_row", create_and_look_from_outside)
    start_time = frozen_clock.current
    await asyncio.wait_for(run_worker(session_factory, storage), timeout=10)  # no deadlock
    assert seen == [start_time, start_time + timedelta(minutes=1)]
    job = admin_client.get(f"{API}/{job_id}").json()
    assert (job["status"], job["created_count"], job["error_count"]) == ("completed", 2, 1)


def test_viewers_cannot_import_but_can_see_history(viewer_client, admin_client):
    assert start(viewer_client, csv_file(row())).status_code == 403
    job_id = start(admin_client, csv_file(row())).json()["id"]
    assert viewer_client.post(f"{API}/{job_id}/commit", json={}).status_code == 403
    assert viewer_client.post(f"{API}/{job_id}/cancel").status_code == 403
    assert viewer_client.get(API).json()["total"] == 1
    assert viewer_client.get(f"{API}/template.xlsx").status_code == 200


def test_unknown_import_is_404(admin_client):
    for path in ("/nope", "/nope/rows", "/nope/errors.csv"):
        assert admin_client.get(API + path).status_code == 404


def test_job_claim_locks_only_the_job_row_on_postgres():
    """Postgres rejects FOR UPDATE on the outer join that loads the job's creator, so the
    lock must name the job table only."""
    sql = str(service.next_job_query().compile(dialect=postgresql.dialect()))
    assert "LEFT OUTER JOIN users" in sql
    assert sql.endswith("FOR UPDATE OF import_jobs SKIP LOCKED")
